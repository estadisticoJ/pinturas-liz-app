"""Convierte el Excel de metadatos de Pinturas Liz (hojas Productos, Clientes y PreciosCliente)
en un archivo SQL para cargar en Supabase y un informe de observaciones.

Reglas de la carga (versión 0.4.1):
- Precios por docena y por cliente. No hay precios de lista.
- Precio repetido con valores distintos: la fila de arriba es el precio de 2025 (cerrado)
  y la de abajo el precio vigente desde 2026. Repetidos con el mismo valor se cargan una vez.
- Precios en cero o vacíos no se cargan: quedan como alertas para los socios.
- Productos sin código ni color se cargan con el color "Único".
- Cada problema encontrado queda como alerta (comentario) para los socios en la aplicación.

Uso:  python scripts/metadatos_a_sql.py Metadatos.xlsx
Crea: supabase/datos/02_carga_metadatos.sql y supabase/datos/observaciones_metadatos.md
Requiere: pip install openpyxl
"""
import sys, re, collections, datetime, pathlib, unicodedata
import openpyxl

ORIGEN = sys.argv[1] if len(sys.argv) > 1 else 'Metadatos.xlsx'
SALIDA = pathlib.Path(__file__).resolve().parent.parent / 'supabase' / 'datos'
SALIDA.mkdir(parents=True, exist_ok=True)

ACENTOS = {'TEMPERA':'Témpera','NEON':'Neón','METALICO':'Metálico','MARRON':'Marrón','LIMON':'Limón','CERULEO':'Cerúleo',
  'VIOLACEO':'Violáceo','ORQUIDEA':'Orquídea','CERAMICO':'Cerámico','SINTETICA':'Sintética','RIGIDO':'Rígido','UNICO':'Único',
  'CARMIN':'Carmín','PURPURA':'Púrpura','LILA':'Lila','MAIZ':'Maíz','FUCSIA':'Fucsia','ACRILICO':'Acrílico','BOTELLON':'Botellón'}
SIGLAS = {'FLT', 'DNI', 'RUC', 'II', 'III'}
MINUSCULAS = {'DE', 'DEL', 'LA', 'LAS', 'LOS', 'Y'}

def clave(v):
    """Nombre para comparar: mayúsculas y sin tildes (TÉMPERA = TEMPERA)."""
    return unicodedata.normalize('NFD', limpio(v).upper()).encode('ascii', 'ignore').decode()

def limpio(v):
    return re.sub(r'\s+', ' ', str(v)).strip() if v is not None else ''

def bonito(texto):
    palabras = limpio(texto).upper().split(' ')
    out = []
    for i, p in enumerate(palabras):
        if p in SIGLAS: out.append(p)
        elif p in ACENTOS: out.append(ACENTOS[p])
        elif i > 0 and p in MINUSCULAS: out.append(p.lower())
        else: out.append(p.capitalize())
    return ' '.join(out)

def q(v):
    return 'null' if v is None or v == '' else "'" + str(v).replace("'", "''") + "'"

# ---------- Línea de producto ----------
def linea(prod):
    p = prod.upper()
    if p.startswith('ALTO RELIEVE'): return 'Alto relieve'
    if p.startswith('TEXTIL') or p.startswith('SUBLIMADO TEXTIL'): return 'Textil'
    if p.startswith('TEMPERA'): return 'Témpera'
    if p.split(' ')[0] in ('MATE', 'BRILLO', 'COUNTRY', 'PERLADO', 'MAGIC'): return 'Acrílicos'
    if 'DIALE' in p or 'DORCO' in p: return 'Otras marcas'
    return 'Complementos y efectos'

# ---------- Color aproximado para las muestras en pantalla ----------
BASE_HEX = [
  ('BLANCO HUESO', '#EFE8D8'), ('BLANCO', '#F7F7F2'), ('MARFIL', '#F3EBD3'), ('PIEL OSCURA', '#C68B65'), ('PIEL', '#F1C6A7'),
  ('AMARILLO LIMON', '#F4EA3C'), ('AMARILLO CADMIO', '#F6B800'), ('AMARILLO CANARIO', '#FFE135'), ('AMARILLO POOH', '#F5B829'), ('AMARILLO', '#F6C900'),
  ('MANDARINA', '#F58A1F'), ('NARANJA', '#F27C1E'), ('MELON', '#F7A576'), ('CORAL', '#F27B6B'),
  ('ROJO INDIO', '#B03A2E'), ('ROJO TEJA', '#B5543A'), ('ROJO VINO', '#7A1E2B'), ('ROJO CARMIN', '#C2183B'), ('ROJO FUEGO', '#E32B1C'), ('ROJO TORO', '#C8102E'),
  ('ROJO', '#D62828'), ('GUINDA', '#7B1C2E'), ('FUCSIA', '#E3157A'), ('MAGENTA', '#C8187E'), ('ROSA', '#F08DB4'), ('ROSADO FUERTE', '#E8569A'), ('ROSADO', '#F08DB4'),
  ('TUTTI', '#F07FB2'), ('BARNEY', '#7E3F98'), ('ORQUIDEA', '#B66DB8'), ('MORADO', '#5E2A84'), ('VIOLETA', '#7B3FA3'), ('VIOLACEO', '#8C4F9E'), ('LILA', '#B9A0D6'),
  ('AZUL NOCHE', '#1B2A55'), ('AZUL ULTRAMAR', '#2338A8'), ('AZUL', '#1F4FBF'), ('CERULEO', '#2A7AB8'), ('CELESTE TURQUESA', '#3BC4CF'), ('CELESTE', '#69BFE8'),
  ('TURQUESA', '#1FB5B0'), ('CALIPSO', '#11A3A6'), ('AGUA MARINA', '#5CCBC0'), ('VERDE AGUA', '#9FDDCB'),
  ('VERDE LIMON', '#9BCB3B'), ('VERDE ESMERALDA', '#1B9E5A'), ('VERDE MANZANA', '#7BC043'), ('VERDE HOJA', '#3C8D3F'), ('VERDE OLIVO', '#6B7A2A'), ('VERDE', '#2E9E4F'),
  ('TIERRA QUEMADA', '#7C3F23'), ('SIENA', '#A0522D'), ('OCRE', '#C78A2C'), ('LADRILLO', '#A4462F'), ('ARCILLA', '#B5734E'), ('CERAMICO', '#B66A45'),
  ('CARAMELO', '#B8742F'), ('TOFFE', '#9C6B3C'), ('NOGAL', '#5C3A21'), ('MARRON', '#7A4A2B'), ('SEPIA', '#704214'), ('MAIZ', '#F2C649'), ('MOSTAZA', '#D2A21B'),
  ('DORADO', '#C9A227'), ('ORO', '#C9A227'), ('PLATEADO', '#B5B9C2'), ('PLATA', '#B5B9C2'), ('COBRE', '#B87333'), ('BRONCE', '#A97142'),
  ('GRIS', '#8E9196'), ('NEGRO', '#1C1C1E'), ('TRANSPARENTE', '#E8EEF2'), ('UNICO', '#D9DCE3'),
]
def mezcla(hexa, hacia, f):
    a = [int(hexa[i:i+2], 16) for i in (1, 3, 5)]; b = [int(hacia[i:i+2], 16) for i in (1, 3, 5)]
    return '#' + ''.join(f'{round(x + (y - x) * f):02X}' for x, y in zip(a, b))
def hex_de(color):
    c = color.upper()
    base = next((h for k, h in BASE_HEX if k in c), '#B9BCC6')
    if 'CLARO' in c or re.search(r'(ITO|ITA)\b', c): base = mezcla(base, '#FFFFFF', 0.45)
    if 'OSCURO' in c: base = mezcla(base, '#000000', 0.3)
    if re.search(r'(ITO|ITA)\b', c) and base == mezcla('#B9BCC6', '#FFFFFF', 0.45):
        for k, h in BASE_HEX:
            raiz = k[:4]
            if c.startswith(raiz): base = mezcla(h, '#FFFFFF', 0.45); break
    return base

# ---------- Lectura ----------
DESDE_ANTERIOR, HASTA_ANTERIOR, DESDE_VIGENTE = '2025-01-01', '2025-12-31', '2026-01-01'

wb = openpyxl.load_workbook(ORIGEN, data_only=True)
productos_x = [r for r in wb['Productos'].iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
clientes_x = [r for r in wb['Clientes'].iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
precios_x = [r for r in wb['PreciosCliente'].iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
obs = collections.defaultdict(list)       # informe en Markdown
alertas = []                              # (tipo, titulo, detalle, cliente_codigo, producto_nombre, medida)

# Productos y colores (la clave de producto no distingue tildes ni mayúsculas)
prod_nombre = {}                          # CLAVE -> nombre bonito
colores = collections.OrderedDict()       # (CLAVE, COLOR) -> (codigo, nombre color)
codigos_usados = set()
for cod, prod, color in productos_x:
    P = clave(prod); C = clave(color) or 'UNICO'
    prod_nombre.setdefault(P, bonito(prod))
    cod = limpio(cod) or None
    if cod and cod in codigos_usados:
        nuevo = cod + 'B'
        obs['Código de producto repetido'].append(f'{cod} se usa en dos colores; a {bonito(prod)} {bonito(C)} se le asignó {nuevo}.')
        alertas.append(('dato_excel', f'Código {cod} repetido', f'{cod} se usa en dos colores. A {bonito(prod)} {bonito(C)} se le asignó {nuevo}; corrígelo si corresponde.', None, P, None))
        cod = nuevo
    if cod: codigos_usados.add(cod)
    if (P, C) in colores:
        obs['Color repetido'].append(f'{bonito(prod)} {bonito(C)} aparece dos veces; se cargó una sola.')
        continue
    colores[(P, C)] = (cod, bonito(C))
    if not cod:
        obs['Producto sin código ni color (se cargó con el color "Único")'].append(bonito(prod))
        alertas.append(('producto_incompleto', f'{bonito(prod)}: sin código ni colores',
            'Está en la hoja Productos sin código ni color. Se cargó con el color "Único" para poder venderlo. Agrega sus colores y códigos cuando los tengas.', None, P, None))

# Clientes
cod_por_nombre = {clave(r[1]): limpio(r[0]) for r in clientes_x}

# Precios: agrupa por cliente, producto y medida en el orden del Excel
grupos = collections.OrderedDict()
for fila, (cli, prod, medida, tipo, precio, vend) in enumerate(precios_x, start=2):
    CL = clave(cli); P = clave(prod); ml = int(medida)
    if P not in prod_nombre:
        prod_nombre[P] = bonito(prod)
        colores[(P, 'UNICO')] = (None, 'Único')
        obs['Producto con precio que no está en la hoja Productos (se creó con color "Único")'].append(bonito(prod))
    if CL not in cod_por_nombre:
        obs['Cliente con precio que no está en la hoja Clientes (no se cargó)'].append(f'Fila {fila}: {bonito(cli)}')
        continue
    if precio is None or precio == '' or float(precio) == 0:
        motivo = 'vacío' if precio is None or precio == '' else 'en cero'
        obs[f'Precio {motivo} (no se cargó; queda como alerta)'].append(f'Fila {fila}: {bonito(cli)} · {prod_nombre[P]} {ml} ml')
        alertas.append(('precio_faltante', f'Falta precio: {prod_nombre[P]} {ml} ml',
            f'En el Excel (fila {fila}) el precio está {motivo}. No se cargó: registra el precio cuando lo tengas.', cod_por_nombre[CL], P, ml))
        continue
    valores = grupos.setdefault((CL, P, ml), [])
    if not valores or valores[-1][0] != float(precio): valores.append((float(precio), fila))
    else: obs['_dup_iguales'].append(fila)

# Cada grupo da uno o más precios con su vigencia
precios = []   # (cliente_codigo, producto, ml, precio, desde, hasta, nota)
for (CL, P, ml), valores in grupos.items():
    cod = cod_por_nombre[CL]
    if len(valores) == 1:
        precios.append((cod, P, ml, valores[0][0], DESDE_VIGENTE, None, 'Carga inicial desde Excel')); continue
    *anteriores, ultimo = valores
    for i, (v, fila) in enumerate(anteriores):
        # varios anteriores: se reparten en meses de 2025, el último cierra el 31/12/2025
        desde = DESDE_ANTERIOR if i == 0 else f'2025-{min(12, 1 + i * 3):02d}-01'
        hasta = HASTA_ANTERIOR if i == len(anteriores) - 1 else (datetime.date.fromisoformat(f'2025-{min(12, 1 + (i + 1) * 3):02d}-01') - datetime.timedelta(days=1)).isoformat()
        precios.append((cod, P, ml, v, desde, hasta, f'Precio anterior (Excel, fila {fila})'))
    precios.append((cod, P, ml, ultimo[0], DESDE_VIGENTE, None, f'Precio vigente (Excel, fila {ultimo[1]})'))
    texto = ' → '.join(f'S/ {v:g}' for v, _ in valores)
    nombre_cli = bonito(next(r[1] for r in clientes_x if clave(r[1]) == CL))
    obs['Precio con valores distintos (se guardó como historial 2025 → 2026)'].append(f'{nombre_cli} · {prod_nombre[P]} {ml} ml: {texto}')
    alertas.append(('precio_revisar', f'Dos precios en el Excel: {prod_nombre[P]} {ml} ml',
        f'Había {texto}. Se guardó S/ {valores[0][0]:g} como precio de 2025 y S/ {ultimo[0]:g} como precio vigente desde 2026. Confirma que el vigente sea el correcto.', cod, P, ml))

con_precio = {p for _, p, *_ in precios}
for P in sorted(prod_nombre, key=lambda x: prod_nombre[x]):
    if P not in con_precio:
        obs['Producto sin ningún precio (no se puede vender hasta que un socio registre medida y precio)'].append(prod_nombre[P])
        alertas.append(('producto_sin_precio', f'{prod_nombre[P]}: sin medida ni precio',
            'No tiene ningún precio en el Excel, por eso no aparece en el catálogo. Un socio debe agregar su medida (ml) y su precio por docena.', None, P, None))

# Presentaciones: cada producto con cada medida que tenga algún precio
presentaciones = sorted({(p, ml) for _, p, ml, *_ in precios})

# Clientes y vendedores
vendedores = sorted({limpio(r[2]).upper() for r in clientes_x if r[2]} | {limpio(r[5]).upper() for r in precios_x if r[5]})
cli_rows = []
for cod, nom, vend, lugar, local, tel, correo, dire in clientes_x:
    cli_rows.append((limpio(cod), bonito(nom), limpio(vend).upper() or None, bonito(lugar) if lugar else None,
                     bonito(local) if local else None, limpio(tel) or None, limpio(correo) or None, limpio(dire) or None))
faltan = [(n, sum(1 for r in cli_rows if not r[i])) for n, i in (('teléfono', 5), ('dirección', 7), ('vendedor', 2))]
for n, k in faltan:
    if k:
        obs['Datos de clientes por completar'].append(f'{k} de {len(cli_rows)} clientes sin {n}')
        alertas.append(('cliente_incompleto', f'{k} clientes sin {n}',
            f'{k} de {len(cli_rows)} clientes no tienen {n}. Complétalo en Clientes, en la ficha de cada uno, a medida que hagan pedidos.', None, None, None))
obs['Datos de clientes por completar'].append(f'Ningún cliente tiene documento (DNI o RUC); hace falta para facturar.')
alertas.append(('cliente_incompleto', 'Clientes sin DNI o RUC', 'Ningún cliente cargado tiene documento. Hace falta para emitir boleta o factura.', None, None, None))

# ---------- SQL ----------
hoy = datetime.date.today().isoformat()
L = []
w = L.append
w('-- Pinturas Liz · Carga de productos, clientes y precios desde Metadatos.xlsx (versión 0.4.1)')
w(f'-- Generado el {hoy} con scripts/metadatos_a_sql.py. Requiere las migraciones 0004 y 0005.')
w('-- Ejecutar una sola vez, después de 01_limpiar_pruebas.sql. Si falla, no queda nada a medias.')
w('')
w('-- Vendedores (luego se vincula cada uno con su usuario en Mi cuenta > Administrar usuarios)')
w('insert into public.vendedores (codigo, nombre) values')
w(',\n'.join(f'  ({q(v)}, {q(bonito(v))})' for v in vendedores))
w('on conflict (codigo) do nothing;')
w('')
w('-- Productos')
w('insert into public.productos (nombre, linea) values')
w(',\n'.join(f'  ({q(n)}, {q(linea(P))})' for P, n in sorted(prod_nombre.items(), key=lambda x: x[1])))
w('on conflict (nombre) do nothing;')
w('')
w('-- Colores (el tono en pantalla es aproximado; se puede ajustar en Supabase > Table Editor > colores)')
w('insert into public.colores (producto_id, codigo, nombre, hex)')
w('select p.id, x.codigo, x.nombre, x.hex from (values')
w(',\n'.join(f'  ({q(prod_nombre[P])}, {q(cod)}, {q(nom)}, {q(hex_de(C))})' for (P, C), (cod, nom) in colores.items()))
w(') as x(producto, codigo, nombre, hex) join public.productos p on p.nombre = x.producto')
w('on conflict (producto_id, nombre) do nothing;')
w('')
w('-- Presentaciones: medida en ml y frascos por paquete (12 = docena; se puede cambiar por presentación)')
w('insert into public.presentaciones (producto_id, medida_ml, unidades, precio)')
w('select p.id, x.ml, 12, 0 from (values')
w(',\n'.join(f'  ({q(prod_nombre[P])}, {m})' for P, m in presentaciones))
w(') as x(producto, ml) join public.productos p on p.nombre = x.producto')
w('on conflict (producto_id, medida_ml, unidades) do nothing;')
w('')
w('-- Clientes')
w('insert into public.clientes (codigo, nombre, vendedor_codigo, lugar, zona, telefono, correo, direccion) values')
w(',\n'.join('  (' + ', '.join(q(v) for v in r) + ')' for r in cli_rows))
w('on conflict (codigo) do nothing;')
w('')
w(f'-- Precios por cliente, por docena. Vigentes desde {DESDE_VIGENTE}; los repetidos con otro valor quedan como historial de 2025.')
w('insert into public.precios (cliente_id, producto_id, medida_ml, tipo_venta, precio_docena, vigente_desde, vigente_hasta, nota)')
w("select c.id, p.id, x.ml, 'menor', x.precio, x.desde::date, x.hasta::date, x.nota from (values")
w(',\n'.join(f'  ({q(c)}, {q(prod_nombre[p])}, {m}, {v:g}, {q(d)}, {q(h)}, {q(n)})' for c, p, m, v, d, h, n in precios))
w(') as x(cliente, producto, ml, precio, desde, hasta, nota)')
w('join public.clientes c on c.codigo = x.cliente join public.productos p on p.nombre = x.producto;')
w('')
w('-- Alertas para los socios: datos faltantes o dudosos encontrados en el Excel')
w("insert into public.alertas (tipo, titulo, detalle, cliente_id, producto_id, medida_ml, origen)")
w("select x.tipo, x.titulo, x.detalle, c.id, p.id, x.ml, 'carga' from (values")
w(',\n'.join(f'  ({q(t)}, {q(ti)}, {q(d)}, {q(c)}, {q(prod_nombre[p]) if p else "null"}, {m if m else "null::integer"})' for t, ti, d, c, p, m in alertas))
w(') as x(tipo, titulo, detalle, cliente, producto, ml)')
w('left join public.clientes c on c.codigo = x.cliente left join public.productos p on p.nombre = x.producto;')
w('')
w('-- Resumen')
w("select 'productos' as tabla, count(*) from public.productos union all select 'colores', count(*) from public.colores")
w("union all select 'presentaciones', count(*) from public.presentaciones union all select 'clientes', count(*) from public.clientes")
w("union all select 'precios vigentes', count(*) from public.precios where vigente_hasta is null")
w("union all select 'precios en historial', count(*) from public.precios where vigente_hasta is not null")
w("union all select 'alertas para socios', count(*) from public.alertas;")
(SALIDA / '02_carga_metadatos.sql').write_text('\n'.join(L) + '\n', encoding='utf-8')

# ---------- Informe ----------
R = [f'# Observaciones de Metadatos.xlsx', '', f'Generado el {hoy}. Cada punto también aparece en la aplicación como alerta para los socios (campana de la barra superior).', '',
     '| Dato | Cantidad cargada |', '| --- | --- |',
     f'| Productos | {len(prod_nombre)} |', f'| Colores | {len(colores)} |', f'| Presentaciones (producto × medida) | {len(presentaciones)} |',
     f'| Clientes | {len(cli_rows)} |', f'| Vendedores | {len(vendedores)} ({", ".join(bonito(v) for v in vendedores)}) |',
     f'| Precios vigentes (desde {DESDE_VIGENTE}) | {sum(1 for x in precios if not x[5])} |', f'| Precios en historial (2025) | {sum(1 for x in precios if x[5])} |',
     f'| Alertas para socios | {len(alertas)} |', '']
iguales = len(obs.pop('_dup_iguales', []))
if iguales: R += [f'- {iguales} precios estaban repetidos con el mismo valor; se cargó uno solo.', '']
for titulo, items in obs.items():
    R += [f'## {titulo} ({len(items)})', ''] + [f'- {i}' for i in sorted(set(items), key=items.index)] + ['']
(SALIDA / 'observaciones_metadatos.md').write_text('\n'.join(R), encoding='utf-8')
print('Listo:', SALIDA / '02_carga_metadatos.sql')
print('\n'.join(R[:40]))
